from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("data/experimental_corpus/huggingface_bilingual/requetes_test_bilingues.docx")


CORPUS = [
    ("Français", "Juridique et réglementation", [
        ("Article 313-19 AMF — Gouvernance des instruments financiers", "fr/legal/amf_v2_7fe23a191b.txt", [
            "Quelles obligations de gouvernance incombent au distributeur d’instruments financiers ?",
            "Comment le distributeur définit-il le marché cible d’un produit financier ?",
            "Quelles informations le distributeur doit-il obtenir avant de recommander un instrument financier ?",
        ]),
        ("BOFiP — Exonération de TVA liée à l’importation", "fr/legal/bofip_b5523a42f9.txt", [
            "Quelles prestations liées à l’importation peuvent être exonérées de TVA ?",
            "Comment prouver qu’une commission est incluse dans la base d’imposition à l’importation ?",
            "Quels frais accessoires sont pris en compte dans la base de TVA à l’importation ?",
        ]),
        ("Jade — Taxe foncière sur des locaux commerciaux", "fr/legal/jade_1b8ed4aa8e.txt", [
            "Pourquoi la demande de réduction de taxe foncière de la SCI Briançon a-t-elle été rejetée ?",
            "Quel local de comparaison a été retenu pour évaluer les locaux commerciaux ?",
            "Quel tarif au mètre carré a été appliqué dans cette affaire ?",
        ]),
        ("Kali — Adhésion à un accord d’intéressement", "fr/legal/kali_3fa8599ccc.txt", [
            "Comment une entreprise de moins de 50 salariés peut-elle adhérer au dispositif d’intéressement ?",
            "Quelles formalités suivent l’adhésion à l’accord d’intéressement ?",
            "Quelle est la durée de l’adhésion et comment fonctionne son renouvellement ?",
        ]),
    ]),
    ("Français", "Actualités", [
        ("FrenchQA — Accord et contrôles administratifs", "fr/news/frenchqa_actualite_1.txt", [
            "Quelles conséquences administratives l’accord ne permet-il pas d’éviter ?",
            "L’accord supprime-t-il les contrôles douaniers ?",
            "Quels obstacles concernent les entreprises qui échangent sans freins ?",
        ]),
        ("FrenchQA — Restrictions pour les éleveurs de visons", "fr/news/frenchqa_actualite_2.txt", [
            "À quelle condition les restrictions imposées aux éleveurs de visons sont-elles assouplies ?",
            "Quel ministère annonce l’assouplissement des mesures sanitaires ?",
            "Que se passe-t-il en l’absence de nouvelles contaminations ?",
        ]),
        ("FrenchQA — Patrimoine minier et tourisme", "fr/news/frenchqa_actualite_3.txt", [
            "Quel organisme a reconnu ce territoire comme patrimoine mondial en 2012 ?",
            "Comment l’ancien pays minier s’est-il réinventé ?",
            "Quel musée a contribué au développement du tourisme local ?",
        ]),
        ("FrenchQA — Risques sanitaires de produits de consommation", "fr/news/frenchqa_actualite_4.txt", [
            "Quels risques sanitaires sont associés à la consommation de ces produits ?",
            "Pourquoi les bénéfices de ces produits sont-ils contestés ?",
            "Quels problèmes de santé ces produits peuvent-ils provoquer ?",
        ]),
    ]),
    ("Anglais", "International news", [
        ("AG News — Venezuelan referendum", "en/news/ag_news_0_1_actualites-internationales.txt", [
            "What was the purpose of the referendum in Venezuela?",
            "Could the vote remove Hugo Chavez from office?",
            "How long would Chavez govern if he won a new mandate?",
        ]),
        ("AG News — Seoul protest over Iraq deployment", "en/news/ag_news_0_2_actualites-internationales.txt", [
            "Why did protesters clash with police in Seoul?",
            "What decision did the South Korean government make about Iraq?",
            "How did police disperse the demonstrators?",
        ]),
    ]),
    ("Anglais", "Sports", [
        ("AG News — Phelps and the 200-meter freestyle", "en/news/ag_news_1_1_sport.txt", [
            "Who qualified for the Olympic 200-meter freestyle semifinals?",
            "Which relay final was Michael Phelps added to?",
            "Who were faster than Phelps in the preliminary heats?",
        ]),
        ("AG News — Reds versus Padres", "en/news/ag_news_1_2_sport.txt", [
            "Which player hit two home runs for the Cincinnati Reds?",
            "What was the score of the Reds-Padres game?",
            "How did the loss affect the Padres' wild-card lead?",
        ]),
    ]),
    ("Anglais", "Economy and business", [
        ("AG News — Wall Street short sellers", "en/news/ag_news_2_1_economie-et-entreprises.txt", [
            "Why are short sellers on Wall Street seeing green again?",
            "What does the article say about Wall Street bears?",
            "Who are described as ultra-cynics in the market?",
        ]),
        ("AG News — Carlyle Group and aerospace", "en/news/ag_news_2_2_economie-et-entreprises.txt", [
            "What sector is Carlyle Group investing in?",
            "Which industry was Carlyle previously known to target?",
            "What new market did Carlyle place bets on?",
        ]),
    ]),
    ("Anglais", "Science and technology", [
        ("AG News — Madden NFL 2005", "en/news/ag_news_3_1_sciences-et-technologies.txt", [
            "Why might office absenteeism have risen on Tuesday?",
            "Which football video game was released?",
            "Why did some fans take a sick day?",
        ]),
        ("AG News — Wireless networking standard", "en/news/ag_news_3_2_sciences-et-technologies.txt", [
            "Which companies proposed a new wireless networking standard?",
            "How much faster is the proposed wireless format?",
            "What kind of networking proposal did technology companies announce?",
        ]),
    ]),
    ("Anglais", "Health and biomedical research", [
        ("PubMedQA — Mitochondria and lace plant cell death", "en/health/pubmedqa_21645374.txt", [
            "Do mitochondria influence programmed cell death in lace plant leaves?",
            "What effect did cyclosporine A have on leaf perforations?",
            "How did mitochondrial dynamics change during programmed cell death?",
        ]),
        ("PubMedQA — Landolt C and Snellen E visual acuity", "en/health/pubmedqa_16418930.txt", [
            "How do Snellen E and Landolt C acuity measurements compare?",
            "Was Snellen E higher than Landolt C in strabismus amblyopia?",
            "What population was tested in the visual acuity study?",
        ]),
        ("PubMedQA — Breast screening in Sami and non-Sami women", "en/health/pubmedqa_22564465.txt", [
            "How did breast cancer screening attendance differ between Sami and non-Sami women?",
            "What were the recall rates in Sami versus non-Sami populations?",
            "Was invasive cancer detection significantly lower in the Sami group?",
        ]),
        ("PubMedQA — Prognosis in alcoholic cirrhosis", "en/health/pubmedqa_23588461.txt", [
            "Can ascites volume predict outcomes in alcoholic cirrhosis?",
            "Which measurements were associated with rehospitalization?",
            "What was the relationship between dry BMI and mortality?",
        ]),
    ]),
]


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in [
        ("Heading 1", 15, "2E74B5", 12, 6),
        ("Heading 2", 12, "2E74B5", 10, 4),
        ("Heading 3", 11, "1F4D78", 7, 3),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    try:
        query_style = doc.styles.add_style("Query item", WD_STYLE_TYPE.PARAGRAPH)
    except ValueError:
        query_style = doc.styles["Query item"]
    query_style.base_style = doc.styles["Normal"]
    query_style.paragraph_format.left_indent = Inches(0.25)
    query_style.paragraph_format.first_line_indent = Inches(-0.25)
    query_style.paragraph_format.space_after = Pt(2)
    query_style.paragraph_format.line_spacing = 1.15


def setup_page(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.60)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.22)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("Projet de recherche vectorielle  |  Guide de requêtes")
    set_run_font(header_run, size=8.5, color=(100, 100, 100))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Page ")
    set_run_font(footer_run, size=8.5, color=(100, 100, 100))
    add_page_field(footer)


def add_document_section(doc, title, path, questions):
    if title.startswith("PubMedQA — Breast screening"):
        continuation = doc.add_paragraph(style="Heading 2")
        continuation.paragraph_format.page_break_before = True
        continuation.paragraph_format.keep_with_next = True
        continuation.add_run("Health and biomedical research (suite)")
    title_p = doc.add_paragraph(style="Heading 3")
    title_p.paragraph_format.keep_with_next = True
    title_p.add_run(title)
    source_p = doc.add_paragraph()
    source_p.paragraph_format.space_after = Pt(3)
    source_p.paragraph_format.keep_with_next = True
    source_run = source_p.add_run(f"Document source : {path}")
    set_run_font(source_run, size=9, color=(102, 102, 102), italic=True)
    for index, question in enumerate(questions, 1):
        paragraph = doc.add_paragraph(style="Query item")
        paragraph.paragraph_format.keep_together = True
        label = paragraph.add_run(f"Q{index}. ")
        set_run_font(label, bold=True, color=(31, 77, 120))
        paragraph.add_run(question)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)
    setup_page(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(22)
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("Guide de requêtes de test")
    set_run_font(title_run, size=24, bold=True, color=(31, 77, 120))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle.add_run("Corpus bilingue de documents expérimentaux")
    set_run_font(subtitle_run, size=13, color=(70, 70, 70))

    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(8)
    note.paragraph_format.left_indent = Inches(0.15)
    note.paragraph_format.right_indent = Inches(0.15)
    shade_paragraph(note, "E8EEF5")
    label = note.add_run("Utilisation. ")
    set_run_font(label, bold=True, color=(31, 77, 120))
    note.add_run(
        "Chaque formulation ci-dessous peut être saisie telle quelle dans l’interface Streamlit ou envoyée à l’API via la route /search. "
        "Les requêtes sont classées selon la langue, le domaine et le document du corpus initial."
    )

    current_language = None
    for language, domain, entries in CORPUS:
        if language != current_language:
            language_heading = doc.add_paragraph(style="Heading 1")
            language_heading.paragraph_format.keep_with_next = True
            language_heading.add_run(language)
            current_language = language
        domain_heading = doc.add_paragraph(style="Heading 2")
        domain_heading.paragraph_format.keep_with_next = True
        domain_heading.add_run(domain)
        for title_text, path, questions in entries:
            add_document_section(doc, title_text, path, questions)

    doc.core_properties.title = "Guide de requêtes de test — corpus bilingue"
    doc.core_properties.subject = "Requêtes de recherche sémantique pour le corpus expérimental"
    doc.core_properties.author = "Projet de recherche vectorielle"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
