from src.extraction import extract_text
from docx import Document

def test_extract_text_from_txt(tmp_path):
    document = tmp_path / "document.txt"
    document.write_text("Bonjour monde", encoding="utf-8")
    assert extract_text(document) == "Bonjour monde"


def test_extract_text_from_docx_includes_tables(tmp_path):
    path = tmp_path / "tableau.docx"
    document = Document()
    document.add_paragraph("Introduction")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Nom"
    table.cell(0, 1).text = "Valeur"
    document.save(path)

    assert extract_text(path).splitlines() == ["Introduction", "Nom | Valeur"]
