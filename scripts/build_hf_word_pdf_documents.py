"""Produit des fichiers Word et PDF importables à partir du corpus Hugging Face."""

from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CORPUS_DIR = PROJECT_ROOT / "data" / "experimental_corpus" / "huggingface_bilingual"
OUTPUT_DIR = CORPUS_DIR / "formatted_documents"
SOFFICE = Path("/Users/joelkouraogo/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice")

SELECTION = [
    "fr/legal/amf_v2_7fe23a191b.txt",
    "fr/news/frenchqa_actualite_3.txt",
    "en/health/pubmedqa_21645374.txt",
    "en/news/ag_news_3_1_sciences-et-technologies.txt",
]


def paragraph_spacing(paragraph, before: int = 0, after: int = 120) -> None:
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(before / 20)
    paragraph_format.space_after = Pt(after / 20)
    paragraph_format.line_spacing = 1.15


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Corpus expérimental bilingue — Hugging Face")
    footer.runs[0].font.size = Pt(8)


def create_docx(entry: dict) -> Path:
    source_path = CORPUS_DIR / entry["file"]
    document = Document()
    configure_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(entry["title"])
    title_run.bold = True
    title_run.font.size = Pt(18)
    paragraph_spacing(title, after=200)

    details = document.add_paragraph()
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    details.add_run(f"Langue : {entry['language']} | Domaine : {entry['domain']}\n").italic = True
    details.add_run(f"Type : {entry['document_type']}").italic = True
    for run in details.runs:
        run.font.size = Pt(10)
    paragraph_spacing(details, after=240)

    heading = document.add_paragraph()
    heading_run = heading.add_run("Contenu du document")
    heading_run.bold = True
    heading_run.font.size = Pt(14)
    paragraph_spacing(heading, before=80, after=120)

    for block in source_path.read_text(encoding="utf-8").split("\n\n"):
        paragraph = document.add_paragraph(block.replace("\\", " ").strip())
        paragraph_spacing(paragraph)

    provenance = document.add_paragraph()
    provenance.add_run("Provenance : ").bold = True
    source = entry["source"]
    provenance.add_run(
        f"jeu de données Hugging Face {source['dataset']} ; split {source['split']} ; "
        f"fichier source {source['file']} ; licence déclarée : {source['license']}."
    )
    for run in provenance.runs:
        run.font.size = Pt(9)
    paragraph_spacing(provenance, before=160, after=0)

    output_path = OUTPUT_DIR / f"{Path(entry['file']).stem}.docx"
    document.save(output_path)
    return output_path


def main() -> None:
    manifest = json.loads((CORPUS_DIR / "manifest.json").read_text(encoding="utf-8"))
    by_file = {entry["file"]: entry for entry in manifest["entries"]}
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for existing in OUTPUT_DIR.glob("*"):
        if existing.is_file():
            existing.unlink()

    docx_files = [create_docx(by_file[file]) for file in SELECTION]
    subprocess.run(
        [str(SOFFICE), "--headless", "--convert-to", "pdf", "--outdir", str(OUTPUT_DIR), *map(str, docx_files)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    print(f"{len(docx_files)} fichiers Word et {len(docx_files)} PDF créés dans {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
